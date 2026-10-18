#!/usr/bin/env python3
"""Build the native PowerPoint version of the ClaimTrace showcase deck.

    python build_pptx.py

The layout mirrors slides.html at 16:9, but the type is sized for projection
rather than for a dense screen page. Every figure on the evidence slide is
copied from docs/accuracy-baseline/ and the test runs at the recorded commit.
"""

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Emu, Inches, Pt

HERE = Path(__file__).resolve().parent
OUT = HERE / "ClaimTrace-Showcase.pptx"

# ---- palette (same tokens as theme.css) -------------------------------------
INK = RGBColor(0x16, 0x18, 0x1D)
INK_SOFT = RGBColor(0x3D, 0x42, 0x4C)
MUTED = RGBColor(0x65, 0x6B, 0x76)
PAPER = RGBColor(0xFB, 0xFB, 0xF9)
PAPER_2 = RGBColor(0xF2, 0xF1, 0xEC)
RULE = RGBColor(0xD9, 0xD7, 0xCF)
RULE_STRONG = RGBColor(0xB9, 0xB6, 0xAB)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)

VERIFIED = RGBColor(0x14, 0x66, 0x3F)
VERIFIED_BG = RGBColor(0xE6, 0xF1, 0xEA)
MISMATCH = RGBColor(0x9A, 0x61, 0x09)
MISMATCH_BG = RGBColor(0xF7, 0xED, 0xDB)
REVIEW = RGBColor(0x6B, 0x5A, 0x9E)
REVIEW_BG = RGBColor(0xEE, 0xEB, 0xF7)
MISSING = RGBColor(0x9B, 0x2F, 0x2A)
MISSING_BG = RGBColor(0xF8, 0xE8, 0xE6)
FAILED = RGBColor(0x4C, 0x4F, 0x57)
FAILED_BG = RGBColor(0xEC, 0xEC, 0xEE)

SERIF = "Georgia"
SANS = "Arial"
MONO = "Courier New"

W, H = 13.333, 7.5
LEFT = 0.8
RIGHT = W - 0.8


def rgb_text(box, text, *, size, font=SANS, color=INK, bold=False, italic=False,
             align=PP_ALIGN.LEFT, spacing=None, line_spacing=1.0):
    tf = box.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = MSO_ANCHOR.TOP
    lines = text.split("\n")
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.line_spacing = line_spacing
        p.space_before = Pt(0)
        p.space_after = Pt(0)
        run = p.add_run()
        run.text = line
        run.font.size = Pt(size)
        run.font.name = font
        run.font.bold = bold
        run.font.italic = italic
        run.font.color.rgb = color
        if spacing is not None:
            run.font._rPr.set("spc", str(spacing))
    return tf


def text(slide, x, y, w, h, value, **kw):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    rgb_text(box, value, **kw)
    return box


def rect(slide, x, y, w, h, fill, line=None, kind=MSO_SHAPE.RECTANGLE, radius=None):
    shape = slide.shapes.add_shape(kind, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    if line is None:
        shape.line.fill.background()
    else:
        shape.line.color.rgb = line
        shape.line.width = Pt(0.75)
    shape.shadow.inherit = False
    if radius is not None and kind == MSO_SHAPE.ROUNDED_RECTANGLE:
        shape.adjustments[0] = radius
    tf = shape.text_frame
    tf.margin_left = tf.margin_right = Inches(0.1)
    tf.margin_top = tf.margin_bottom = 0
    tf.word_wrap = False
    return shape


def rule(slide, x1, y1, x2, y2, color=RULE, width=0.75):
    line = slide.shapes.add_connector(
        MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2)
    )
    line.line.color.rgb = color
    line.line.width = Pt(width)
    return line


def chip(slide, x, y, label, fg, bg, width):
    shape = rect(slide, x, y, width, 0.27, bg, kind=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.16)
    tf = shape.text_frame
    tf.margin_left = tf.margin_right = 0
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    run = p.add_run()
    run.text = "●  " + label.upper()
    run.font.size = Pt(8)
    run.font.name = MONO
    run.font.color.rgb = fg
    run.font._rPr.set("spc", "40")
    return shape


def eyebrow(slide, label, right=""):
    text(slide, LEFT, 0.40, 8.0, 0.25, label.upper(), size=9, font=MONO,
         color=MUTED, spacing=120)
    if right:
        text(slide, RIGHT - 5.0, 0.40, 5.0, 0.25, right, size=9, font=MONO,
             color=RULE_STRONG, align=PP_ALIGN.RIGHT, spacing=60)
    rule(slide, LEFT, 0.72, RIGHT, 0.72)


def headline(slide, value, y=0.95, size=31, w=11.8, color=INK, italic_tail=None):
    box = slide.shapes.add_textbox(Inches(LEFT), Inches(y), Inches(w), Inches(1.4))
    tf = box.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    lines = value.split("\n")
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.line_spacing = 1.05
        p.space_after = Pt(0)
        run = p.add_run()
        run.text = line
        run.font.size = Pt(size)
        run.font.name = SERIF
        run.font.color.rgb = color
    return box


def notes(slide, value):
    slide.notes_slide.notes_text_frame.text = value


def base(prs, fill=PAPER):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    bg = slide.background.fill
    bg.solid()
    bg.fore_color.rgb = fill
    return slide


def build():
    prs = Presentation()
    prs.slide_width = Inches(W)
    prs.slide_height = Inches(H)

    # ---------------------------------------------------------------- 01 title
    s = base(prs)
    text(s, 0.9, 2.10, 8.0, 0.3, "SOFTWARE INNOVATION STUDIO · GROUP 18",
         size=10, font=MONO, color=VERIFIED, spacing=140)
    text(s, 0.85, 2.38, 9.0, 1.5, "ClaimTrace", size=82, font=SERIF, color=INK)
    text(s, 0.9, 3.92, 8.4, 0.7, "Check what your citations actually say.",
         size=25, font=SERIF, italic=True, color=INK_SOFT)
    rule(s, 0.9, 5.32, 8.6, 5.32)
    text(s, 0.9, 5.48, 8.2, 0.3,
         "WEB WORKSPACE   ·   OVERLEAF EXTENSION   ·   PDF PARSER   ·   VERIFICATION ENGINE",
         size=8.5, font=MONO, color=MUTED, spacing=50)

    text(s, 9.65, 1.55, 2.9, 0.28, "[P2]", size=9, font=MONO, color=RULE_STRONG,
         align=PP_ALIGN.RIGHT, spacing=100)
    text(s, 9.65, 1.86, 2.9, 1.5,
         "“I want the evidence, not just the conclusion — so I can tell the student "
         "exactly what to fix.”",
         size=12.5, font=SERIF, italic=True, color=INK_SOFT, align=PP_ALIGN.RIGHT,
         line_spacing=1.25)
    text(s, 9.65, 3.55, 2.9, 0.7,
         "— supervisor persona · translated from user research",
         size=8, font=MONO, color=MUTED, align=PP_ALIGN.RIGHT, line_spacing=1.3)
    notes(s, "30 seconds. ClaimTrace checks what a citation actually says. "
             "The tagline is the product: not a score, not a label — the evidence.")

    # -------------------------------------------------------------- 02 problem
    s = base(prs)
    eyebrow(s, "02 · The problem", "ClaimTrace")
    h = s.shapes.add_textbox(Inches(LEFT), Inches(0.95), Inches(11.8), Inches(1.5))
    tf = h.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    p = tf.paragraphs[0]
    p.line_spacing = 1.05
    for value, color, italic in [
        ("The last 48 hours go into checking citations, ", INK, False),
        ("not writing the paper.", VERIFIED, True),
    ]:
        r = p.add_run()
        r.text = value
        r.font.size = Pt(31)
        r.font.name = SERIF
        r.font.italic = italic
        r.font.color.rgb = color

    cards = [
        ("UNTRUSTED", "Co-authors’ references are taken on faith.",
         "A Related Work section written by three people contains references nobody has opened."),
        ("INTERRUPTED", "Every check means leaving the editor.",
         "Overleaf → Zotero → a PDF reader → back. The cost is not the click, it is the context switch."),
        ("FABRICATED", "Machine-written references look real.",
         "Wrong venue, wrong year, or a DOI that resolves to a different paper entirely."),
    ]
    for i, (kicker, title, body) in enumerate(cards):
        x = LEFT + i * 3.95
        rule(s, x, 2.62, x + 3.55, 2.62, INK, 1.5)
        text(s, x, 2.78, 3.55, 0.25, kicker, size=8.5, font=MONO, color=VERIFIED, spacing=120)
        text(s, x, 3.08, 3.55, 0.85, title, size=17, font=SERIF, color=INK, line_spacing=1.1)
        text(s, x, 4.00, 3.55, 1.1, body, size=11.5, font=SANS, color=MUTED, line_spacing=1.35)

    rect(s, LEFT, 6.05, 0.045, 0.85, VERIFIED)
    text(s, LEFT + 0.25, 6.10, 9.4, 0.8,
         "All three personas asked for the same thing:\nthe original evidence, not just a conclusion.",
         size=15, font=SERIF, italic=True, color=INK, line_spacing=1.25)
    text(s, 10.55, 6.18, 2.0, 0.4, "user-research/personas.md", size=8.5, font=MONO,
         color=MUTED, align=PP_ALIGN.RIGHT)
    notes(s, "45 seconds. Three failure modes, one cause: checking is expensive, so it gets skipped. "
             "Every persona we interviewed asked for the source, not a verdict.")

    # ------------------------------------------------------------- 03 pillars
    s = base(prs)
    eyebrow(s, "03 · What it does", "Three surfaces, one evidence trail")
    headline(s, "Three checks, in the place\nyou already write.")
    pillars = [
        ("01 / AUDIT", "Audit", "every reference, every field",
         "Each bibliography entry is checked against OpenAlex, then Crossref. "
         "Title, authors, year and venue are compared field by field.",
         ["Identifier first, title only as fallback",
          "Five states — never a silent “not found”",
          "Works on .bib files and whole PDFs"]),
        ("02 / VERIFY", "Verify", "claim vs. what the paper says",
         "A selected sentence is matched to the cited paper, the relevant passages "
         "are retrieved, and the judgement ships with its evidence.",
         ["Four verdicts, every one evidenced",
          "Every verdict carries the passage",
          "Bring your own model — no server key"]),
        ("03 / READ", "Read", "the paper as it was written",
         "PDFs become structured Markdown — formulas, tables and figures preserved — "
         "and render inside the review pane.",
         ["Hybrid parsing, local fallback",
          "Page-anchored paragraphs",
          "Hover a citation in Overleaf"]),
    ]
    for i, (n, title, sub, body, bullets) in enumerate(pillars):
        x = LEFT + i * 3.95
        rect(s, x, 2.72, 3.62, 3.42, PAPER_2)
        text(s, x + 0.25, 2.95, 3.1, 0.25, n, size=8.5, font=MONO, color=MUTED, spacing=100)
        text(s, x + 0.25, 3.20, 3.1, 0.55, title, size=27, font=SERIF, color=INK)
        text(s, x + 0.25, 3.72, 3.1, 0.25, sub, size=8.5, font=MONO, color=VERIFIED, spacing=30)
        text(s, x + 0.25, 4.05, 3.1, 0.95, body, size=10.5, font=SANS, color=INK_SOFT,
             line_spacing=1.3)
        text(s, x + 0.25, 5.10, 3.1, 1.0, "—  " + "\n—  ".join(bullets),
             size=8.5, font=SANS, color=MUTED, line_spacing=1.4)
    notes(s, "60 seconds. Three surfaces: audit the bibliography, verify a claim against "
             "the cited passage, and read the paper inside the tool. Read is what makes "
             "the other two cheap — no download-and-Ctrl+F loop.")

    # --------------------------------------------------------------- 04 how
    s = base(prs)
    eyebrow(s, "04 · How it works", "One process, two outbound boundaries")
    headline(s, "The engine decides identity.\nThe product never guesses.")
    # clients
    for y, title, body in [(2.85, "CLIENT", "React web workspace"),
                           (3.88, "CLIENT", "Overleaf extension\nManifest V3 · hover")]:
        rect(s, LEFT, y, 2.45, 0.86, WHITE, RULE_STRONG)
        text(s, LEFT + 0.18, y + 0.14, 2.1, 0.2, title, size=8, font=MONO, color=MUTED, spacing=100)
        text(s, LEFT + 0.18, y + 0.38, 2.1, 0.45, body, size=11, font=SANS, color=INK,
             line_spacing=1.15)
    # backend
    rect(s, 3.62, 2.85, 3.15, 1.00, INK)
    text(s, 3.80, 2.99, 2.8, 0.2, "BACKEND", size=8, font=MONO, color=RULE_STRONG, spacing=100)
    text(s, 3.80, 3.22, 2.8, 0.55, "FastAPI\nREST /api/*", size=11.5, font=SANS, color=WHITE,
         line_spacing=1.15)
    text(s, 3.62, 3.90, 3.15, 0.3, "↓  in-process", size=9, font=MONO, color=RULE_STRONG,
         align=PP_ALIGN.CENTER)
    rect(s, 3.62, 4.22, 3.15, 0.98, WHITE, RULE_STRONG)
    text(s, 3.80, 4.36, 2.8, 0.2, "ENGINE + PARSER", size=8, font=MONO, color=MUTED, spacing=100)
    text(s, 3.80, 4.60, 2.8, 0.5, "Identity rules · retrieval\nPDF → Markdown",
         size=11, font=SANS, color=INK, line_spacing=1.15)
    # outbound
    rect(s, 7.15, 2.85, 5.38, 0.98, WHITE, RULE_STRONG)
    rect(s, 7.15, 2.85, 0.05, 0.98, VERIFIED)
    text(s, 7.38, 2.99, 4.9, 0.2, "OUTBOUND · PUBLIC METADATA", size=8, font=MONO,
         color=MUTED, spacing=100)
    text(s, 7.38, 3.24, 4.9, 0.4, "OpenAlex → Crossref — identifier first, title second",
         size=11.5, font=SANS, color=INK)
    rect(s, 7.15, 4.08, 5.38, 0.98, WHITE, RULE_STRONG)
    rect(s, 7.15, 4.08, 0.05, 0.98, VERIFIED)
    text(s, 7.38, 4.22, 4.9, 0.2, "OUTBOUND · USER-OWNED MODEL", size=8, font=MONO,
         color=MUTED, spacing=100)
    text(s, 7.38, 4.47, 4.9, 0.4, "17 providers · the key travels in the request, never in the server",
         size=11.5, font=SANS, color=INK)

    rect(s, LEFT, 6.02, RIGHT - LEFT, 0.78, VERIFIED_BG)
    text(s, LEFT + 0.25, 6.20, RIGHT - LEFT - 0.5, 0.45,
         "A failed lookup is never reported as “not found”.   Five states in Audit, "
         "and in Verify only COMPARED carries a verdict.",
         size=13.5, font=SERIF, color=VERIFIED)
    notes(s, "45 seconds. Clients only speak REST. The engine owns identity. "
             "There are exactly two outbound boundaries — public metadata, and the "
             "user's own model key. The server never holds an application key.")

    # ------------------------------------------------------------- 05 audit
    s = base(prs)
    eyebrow(s, "05 · Audit", "What the user sees")
    headline(s, "A reference is either checked,\nor it is visibly not checked.")
    y_header = 2.38
    text(s, LEFT, y_header, 3.0, 0.25, "REFERENCE", size=8, font=MONO, color=MUTED, spacing=100)
    text(s, 3.95, y_header, 2.2, 0.25, "STATE", size=8, font=MONO, color=MUTED, spacing=100)
    text(s, 6.35, y_header, 6.0, 0.25, "EVIDENCE", size=8, font=MONO, color=MUTED, spacing=100)
    rule(s, LEFT, 2.66, RIGHT, 2.66, RULE_STRONG, 1.0)

    rows = [
        ("bert-exact", "Devlin et al., 2019",
         "All four fields match the OpenAlex record. One lookup; Crossref is never asked.",
         "title · authors · year · venue → MATCH",
         "Verified", VERIFIED, VERIFIED_BG, 1.05),
        ("bert-year-plus-one", "year stated one late",
         "The record stays eligible inside the one-year tolerance, so the difference is "
         "reported field by field.",
         "year → MISMATCH, and the rest of the record is still checked",
         "Metadata mismatch", MISMATCH, MISMATCH_BG, 1.85),
        ("dialoguernn-year-minus-five", "year stated five early",
         "A five-year gap rejects every record both providers return, so the report carries "
         "field_checks: [] rather than a fabricated mismatch.",
         "no acceptable record → nothing to compare",
         "Not found", MISSING, MISSING_BG, 1.1),
        ("provider-500", "one provider errors",
         "Crossref’s search is negative and OpenAlex cannot answer. The reference stays "
         "visibly unchecked — never reported as “not found”.",
         "an unread reference is not a missing reference",
         "Lookup failed", FAILED, FAILED_BG, 1.3),
    ]
    row_h = 0.80
    for i, (ref, sub, evidence, why, label, fg, bg, chip_w) in enumerate(rows):
        y = 2.74 + i * row_h
        text(s, LEFT, y, 3.0, 0.55, ref, size=9.5, font=MONO, color=INK_SOFT, line_spacing=1.2)
        text(s, LEFT, y + 0.24, 3.0, 0.3, sub, size=9.5, font=MONO, color=MUTED)
        chip(s, 3.95, y + 0.06, label, fg, bg, chip_w)
        text(s, 6.35, y + 0.01, 6.1, 0.52, evidence, size=10, font=SANS, color=INK_SOFT,
             line_spacing=1.16)
        text(s, 6.35, y + 0.55, 6.1, 0.22, why, size=8.5, font=MONO, color=MUTED)
        if i < len(rows) - 1:
            rule(s, LEFT, y + row_h - 0.06, RIGHT, y + row_h - 0.06, RULE)

    y_legend = 6.18
    text(s, LEFT, y_legend + 0.03, 1.0, 0.3, "FIVE STATES", size=8, font=MONO,
         color=MUTED, spacing=100)
    lx = 1.78
    for label, fg, bg, width in [
        ("Verified", VERIFIED, VERIFIED_BG, 1.0),
        ("Metadata mismatch", MISMATCH, MISMATCH_BG, 1.75),
        ("Needs review", REVIEW, REVIEW_BG, 1.25),
        ("Not found", MISSING, MISSING_BG, 1.0),
        ("Lookup failed", FAILED, FAILED_BG, 1.2),
    ]:
        chip(s, lx, y_legend, label, fg, bg, width)
        lx += width + 0.12
    text(s, 9.55, y_legend + 0.03, 3.0, 0.3,
         "“Lookup failed” is a state, not a missing reference.",
         size=8.5, font=SANS, color=MUTED, align=PP_ALIGN.RIGHT)
    notes(s, "60 seconds. This is the core invariant. A record that is rejected before "
             "comparison has no field checks — we never invent a mismatch. A provider that "
             "fails leaves the reference visibly unchecked. Four real fixtures, no mock UI.")

    # ------------------------------------------------------------ 06 verify
    s = base(prs)
    eyebrow(s, "06 · Verify", "Claim × cited source")
    headline(s, "The verdict is only the headline.\nThe passage is the argument.")
    # left panel
    rect(s, LEFT, 2.62, 5.78, 3.42, WHITE, RULE_STRONG)
    text(s, LEFT + 0.28, 2.82, 5.2, 0.22, "CLAIM 01", size=8.5, font=MONO, color=MUTED, spacing=100)
    text(s, LEFT + 0.28, 3.12, 5.2, 1.2,
         "“The paper reports that the method reduced carbon emissions by forty percent.”",
         size=16, font=SERIF, color=INK, line_spacing=1.25)
    chip(s, LEFT + 0.28, 5.48, "Not found", MISSING, MISSING_BG, 1.1)
    text(s, LEFT + 1.55, 5.50, 3.9, 0.5,
         "The source was read and does not contain this claim. That is different from failing to read it.",
         size=8.5, font=SANS, color=MUTED, line_spacing=1.2)

    # right panel
    rect(s, 6.72, 2.62, 5.81, 3.42, PAPER_2)
    text(s, 7.00, 2.82, 5.2, 0.22, "CLAIM 02", size=8.5, font=MONO, color=MUTED, spacing=100)
    text(s, 7.00, 3.12, 5.2, 0.70,
         "“Self-attention relates different positions of a single sequence.”",
         size=16, font=SERIF, color=INK, line_spacing=1.25)
    rect(s, 7.00, 3.86, 0.035, 1.22, VERIFIED)
    text(s, 7.18, 3.86, 5.0, 1.30,
         "“Self-attention, sometimes called intra-attention, is an attention mechanism "
         "relating different positions of a single sequence in order to compute a "
         "representation of the sequence.”",
         size=10.5, font=SANS, color=INK_SOFT, line_spacing=1.25)
    chip(s, 7.00, 5.48, "Support", VERIFIED, VERIFIED_BG, 0.95)
    text(s, 8.08, 5.50, 4.2, 0.3,
         "Retrieved from the cited paper · quoted verbatim.",
         size=8.5, font=SANS, color=MUTED)

    text(s, LEFT, 6.30, 11.9, 0.75,
         "Every comparison ships the retrieved passages, their rank and similarity score, and the page "
         "the passage came from when the parser recorded one. A model that ran is not a model that "
         "judged — COMPARED is the only status that carries a verdict.",
         size=10, font=SANS, color=MUTED, line_spacing=1.3)
    notes(s, "60 seconds. Two claims against the same paper. One is not in the source; "
             "one is supported by a passage you can read. The verdict is never shown without "
             "the passage, and only COMPARED carries one.")

    # ----------------------------------------------------------- 07 evidence
    s = base(prs)
    eyebrow(s, "07 · Evidence", "Measured, not claimed")
    headline(s, "What we can defend with a number.")
    numbers = [
        ("1,130", "AUTOMATED TESTS\nACROSS FIVE SUITES"),
        ("34/34", "CONTROLLED AUDIT CASES\nEXACT · MACRO-F1 1.00"),
        ("12/12", "LIVE VERIFY CASES\nDEEPSEEK · 12 CALLS"),
        ("0", "SERVER-SIDE MODEL KEYS\nBRING YOUR OWN"),
    ]
    for i, (big, label) in enumerate(numbers):
        x = LEFT + i * 3.12
        rule(s, x, 2.28, x + 2.85, 2.28, INK, 1.5)
        text(s, x, 2.42, 2.85, 0.85, big, size=44, font=SERIF, color=INK)
        text(s, x, 3.26, 2.85, 0.5, label, size=8, font=MONO, color=MUTED, line_spacing=1.45)

    # distribution bar
    text(s, LEFT, 4.12, 8.0, 0.22, "WHERE THE 31 CONTROLLED .BIB CASES LAND",
         size=8, font=MONO, color=MUTED, spacing=100)
    text(s, 7.4, 4.12, 5.1, 0.22, "STATUS ACCURACY 1.00 · MACRO-F1 1.00",
         size=8, font=MONO, color=VERIFIED, align=PP_ALIGN.RIGHT, spacing=40)
    bar_x, bar_w, bar_y, bar_h = LEFT, RIGHT - LEFT, 4.40, 0.40
    segments = [(7, VERIFIED), (12, MISMATCH), (9, REVIEW), (1, MISSING), (2, FAILED)]
    cursor = bar_x
    for count, color in segments:
        width = bar_w * count / 31
        rect(s, cursor, bar_y, width, bar_h, color)
        text(s, cursor, bar_y + 0.10, width, 0.22, str(count), size=9, font=MONO,
             color=WHITE, align=PP_ALIGN.CENTER)
        cursor += width

    lx = LEFT
    for label, fg, bg, width in [
        ("Verified 7", VERIFIED, VERIFIED_BG, 1.15),
        ("Metadata mismatch 12", MISMATCH, MISMATCH_BG, 2.05),
        ("Needs review 9", REVIEW, REVIEW_BG, 1.55),
        ("Not found 1", MISSING, MISSING_BG, 1.25),
        ("Lookup failed 2", FAILED, FAILED_BG, 1.45),
    ]:
        chip(s, lx, 4.98, label, fg, bg, width)
        lx += width + 0.15

    rule(s, LEFT, 5.95, RIGHT, 5.95, RULE)
    text(s, LEFT, 6.08, 11.9, 0.95,
         "Audit measured on 31 .bib and 3 PDF fixtures with pinned provider responses "
         "(generated 2026-10-08): state accuracy 1.00, every field check exact. Verify measured "
         "on six manually labelled claims in two modes against a seeded excerpt of Attention Is "
         "All You Need; labels are fixture labels, not independent expert annotation. Test count "
         "at local main e171fb3; overhead is linear at 0.055 s per audit request in-process.",
         size=8.5, font=SANS, color=MUTED, line_spacing=1.35)
    notes(s, "45 seconds. Two kinds of evidence: the whole suite is green, and the controlled "
             "audit set is exact. State the caveat out loud — the verify labels are fixture "
             "labels, which is why the annotation work is the next step.")

    # -------------------------------------------------------------- 08 close
    s = base(prs)
    eyebrow(s, "08 · Close", "Group 18")
    h = s.shapes.add_textbox(Inches(LEFT), Inches(1.45), Inches(11.8), Inches(1.9))
    tf = h.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    p1 = tf.paragraphs[0]
    p1.line_spacing = 1.05
    r = p1.add_run()
    r.text = "Every citation, checked —"
    r.font.size = Pt(38)
    r.font.name = SERIF
    r.font.color.rgb = INK
    p2 = tf.add_paragraph()
    p2.line_spacing = 1.05
    r = p2.add_run()
    r.text = "with the evidence."
    r.font.size = Pt(38)
    r.font.name = SERIF
    r.font.italic = True
    r.font.color.rgb = VERIFIED

    nexts = [
        ("NEXT · 01", "Live verification at corpus scale, with the annotation set the team is still labelling."),
        ("NEXT · 02", "In-editor selection and verdicts without leaving Overleaf."),
        ("NEXT · 03", "Deployment hardening: one worker, file-backed storage, a documented restore path."),
    ]
    for i, (kicker, body) in enumerate(nexts):
        x = LEFT + i * 3.95
        rule(s, x, 3.95, x + 3.55, 3.95, RULE_STRONG)
        text(s, x, 4.10, 3.55, 0.22, kicker, size=8, font=MONO, color=MUTED, spacing=100)
        text(s, x, 4.38, 3.55, 0.9, body, size=11.5, font=SANS, color=INK_SOFT, line_spacing=1.35)

    text(s, LEFT, 6.18, 8.0, 0.3,
         "CLAIMTRACE · SOFTWARE INNOVATION STUDIO 41129 · 2026",
         size=8, font=MONO, color=MUTED, spacing=80)
    notes(s, "30 seconds. Close on the invariant. Then the three next steps: corpus-scale "
             "verification, in-editor verdicts, and deployment hardening.")

    prs.save(OUT)
    return OUT


if __name__ == "__main__":
    path = build()
    print(path)
